"""/export: CSV or Excel (.xlsx) with the same rows."""
import datetime as dt
import io

from openpyxl import load_workbook

from bot.db import crud
from bot.handlers.export import _filename, _fmt
from bot.i18n import i18n
from bot.services.export import build_excel
from tests.conftest import make_receipt

JUNE = dt.date(2026, 6, 15)


async def _data(session):
    await make_receipt(session, store="Café Wien", currency="EUR", total=10.0, total_pln=43.0, date=JUNE,
                       items=[{"name": "Kaffee", "total_price": 6.0, "category": "cafe"},
                              {"name": "Brot", "total_price": 4.0, "category": "groceries"}])
    receipts = await crud.get_transactions_for_export(session, 1, None, None)
    items = await crud.get_items_for_export(session, 1, None, None)
    return receipts, items


async def test_xlsx_has_localized_sheets_real_dates_and_pln(db_session):
    receipts, items = await _data(db_session)
    with i18n.use_locale("pl"):
        wb = load_workbook(io.BytesIO(build_excel(receipts, items)))

    assert wb.sheetnames == ["Transakcje", "Pozycje"]
    tx, it = wb["Transakcje"], wb["Pozycje"]
    assert [c.value for c in tx[1]] == ["Data", "Sklep", "Razem", "Waluta", "Razem PLN", "Kategoria"]
    assert tx["A2"].value == dt.datetime(2026, 6, 15) and tx["A2"].number_format == "DD.MM.YYYY"
    assert (tx["C2"].value, tx["D2"].value, tx["E2"].value) == (10.0, "EUR", 43.0)
    assert tx["E2"].number_format == "#,##0.00"
    assert tx.freeze_panes == "A2" and tx.auto_filter.ref

    rows = [[c.value for c in row] for row in it.iter_rows(min_row=2)]
    kaffee = next(r for r in rows if r[2] == "Kaffee")
    assert kaffee[5:8] == [6.0, "EUR", 25.8]  # native price, currency, PLN at the receipt's rate


async def test_xlsx_only_includes_the_requested_sheet(db_session):
    receipts, items = await _data(db_session)
    with i18n.use_locale("en"):
        assert load_workbook(io.BytesIO(build_excel(receipts=receipts))).sheetnames == ["Transactions"]
        assert load_workbook(io.BytesIO(build_excel(item_rows=items))).sheetnames == ["Items"]


def test_legacy_callbacks_mean_csv_and_filename_follows_format():
    assert _fmt(None) == "csv" and _fmt("bogus") == "csv" and _fmt("xlsx") == "xlsx"
    assert _filename("all", JUNE, JUNE, "xlsx") == "financebot_all_15.06.2026.xlsx"
    assert _filename("items", None, None) == "financebot_items_all.csv"
