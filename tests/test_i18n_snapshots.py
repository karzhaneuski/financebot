"""Golden snapshots of every user-visible message, per language.

Each scenario drives the real dispatcher (or calls a service / API route
directly) against a fixed in-memory dataset at a frozen date, recording
everything that would reach the user: message texts, captions, keyboard
buttons, callback alerts, chart/Wrapped canvas texts, Excel headers and API
labels.

Regenerate after an intentional text change with:
    UPDATE_SNAPSHOTS=1 pytest tests/test_i18n_snapshots.py
and review the JSON diff before committing it.
"""
from __future__ import annotations

import datetime as dt
import json
import os
from pathlib import Path

import fakeredis.aioredis
import pytest
from freezegun import freeze_time

import bot.main  # noqa: F401  -- import before freezegun kicks in (pydantic/fastapi break when imported frozen)
import bot.services.wrapped  # noqa: F401
from bot.db.models import Budget, Category, Item, Receipt, ReportSettings
from tests import i18n_harness as h

SNAPSHOT_DIR = Path(__file__).parent / "snapshots"
FROZEN_NOW = "2026-06-20 12:00:00"
U = 1001        # user with a full dataset
EMPTY = 2002    # user without any data
RESET = 3003    # user whose data gets wiped by /reset confirm


# ── dataset ──────────────────────────────────────────────────────────────────

async def _receipt(session, *, user_id=U, store, day, total, total_pln=None, currency="PLN",
                   photo=None, tx_type="purchase", source=None, category=None, items=()):
    r = Receipt(
        user_id=user_id, store=store, date=day, currency=currency, total=total,
        total_pln=total if total_pln is None else total_pln, photo_file_id=photo,
        tx_type=tx_type, source=source, category=category,
    )
    session.add(r)
    await session.flush()
    for it in items:
        session.add(Item(
            receipt_id=r.id, name=it["name"], normalized_name=it.get("norm"),
            quantity=it.get("qty", 1), unit_price=it["total"] / it.get("qty", 1),
            total_price=it["total"], category=Category(it["cat"]), volume_ml=it.get("vol"),
        ))
    await session.flush()
    return r


async def seed(session) -> dict[str, int]:
    d = dt.date
    ids: dict[str, int] = {}
    r = await _receipt(session, store="Kaufland", day=d(2026, 6, 18), total=120, photo="f1", items=[
        {"name": "CocaCola 0,5L", "norm": "coca-cola", "qty": 2, "total": 12, "cat": "groceries", "vol": 500},
        {"name": "Chleb", "norm": "chleb", "total": 8, "cat": "groceries"},
        {"name": "Szampon", "norm": "szampon", "total": 100, "cat": "household"},
    ])
    ids["kaufland"] = r.id
    await _receipt(session, store="Lidl", day=d(2026, 6, 19), total=60, photo="f2", items=[
        {"name": "Coca-Cola", "norm": "coca-cola", "qty": 3, "total": 18, "cat": "groceries", "vol": 500},
        {"name": "Mleko", "norm": "mleko", "total": 42, "cat": "groceries"},
    ])
    await _receipt(session, store="Starbucks", day=d(2026, 6, 19), total=10, total_pln=43,
                   currency="EUR", photo="f3", items=[{"name": "Latte", "total": 10, "cat": "cafe"}])
    await _receipt(session, store="Снятие наличных", day=d(2026, 6, 10), total=200,
                   tx_type="cash_withdrawal", source="erste",
                   items=[{"name": "Снятие наличных", "total": 200, "cat": "other"}])
    await _receipt(session, store="Stypendium", day=d(2026, 6, 5), total=1500, tx_type="income",
                   source="erste", items=[{"name": "Stypendium", "total": 1500, "cat": "other"}])
    for month in (5, 6):
        await _receipt(session, store="Академик", day=d(2026, month, 1), total=900, source="erste",
                       items=[{"name": "Академик", "total": 900, "cat": "housing"}])
    for month in (3, 4, 5, 6):
        await _receipt(session, store="Spotify", day=d(2026, month, 5), total=23.99, source="revolut",
                       category=Category.subscriptions,
                       items=[{"name": "Spotify", "total": 23.99, "cat": "subscriptions"}])
    await _receipt(session, store=None, day=d(2026, 6, 12), total=25,
                   items=[{"name": "Расход", "total": 25, "cat": "transport"}])
    await _receipt(session, store=None, day=d(2026, 6, 15), total=50, source="erste",
                   items=[{"name": "Оплата картой за рубежом", "total": 50, "cat": "other"}])
    await _receipt(session, store="Kaufland", day=d(2026, 5, 10), total=80, photo="f4", items=[
        {"name": "Coca-Cola", "norm": "coca-cola", "total": 20, "cat": "groceries", "vol": 500},
        {"name": "Chleb", "norm": "chleb", "total": 60, "cat": "groceries"},
    ])
    await _receipt(session, store="Pizzeria", day=d(2026, 5, 20), total=55, photo="f5",
                   items=[{"name": "Pizza", "total": 55, "cat": "cafe"}])
    await _receipt(session, user_id=RESET, store="Biedronka", day=d(2026, 6, 1), total=10,
                   items=[{"name": "Bułka", "total": 10, "cat": "groceries"}])

    for cat, limit in (("groceries", 150), ("cafe", 40), ("transport", 30)):
        session.add(Budget(user_id=U, category=Category(cat), limit_pln=limit, month="2026-06",
                           last_notified_pct=0))
    session.add(ReportSettings(user_id=U, daily_enabled=True, weekly_enabled=True, monthly_enabled=False))
    await session.flush()
    return ids


# ── scenario plumbing ────────────────────────────────────────────────────────

class Ctx:
    def __init__(self, session, redis, bot, dp, ids, monkeypatch):
        self.session, self.redis, self.bot, self.dp = session, redis, bot, dp
        self.ids, self.monkeypatch = ids, monkeypatch

    async def feed(self, update):
        async with h.use_session(self.session):
            await self.dp.feed_update(self.bot, update)

    async def msg(self, text, user_id=U, **kw):
        await self.feed(h.message_update(h.tg_user(user_id), text, **kw))

    async def cb(self, data, user_id=U):
        await self.feed(h.callback_update(h.tg_user(user_id), data))

    def ret(self, value):
        if isinstance(value, str):
            h.RECORDER.add({"m": "return", "text": value})
        else:
            h.RECORDER.add({"m": "return", "value": value})


def _patch_receipt(ctx, **attrs):
    import bot.handlers.receipt as rh

    async def _noop(*a, **k):
        return None
    ctx.monkeypatch.setattr(rh, "normalize_item_names", _noop)
    for name, value in attrs.items():
        ctx.monkeypatch.setattr(rh, name, value)


def _async_return(value=None, exc=None):
    async def _f(*a, **k):
        if exc is not None:
            raise exc
        return value
    return _f


PARSED_RECEIPT = {
    "store": "Żabka", "date": "2026-06-20", "currency": "PLN", "total": 31.5,
    "items": [
        {"name": "Hot-dog", "quantity": 2, "unit_price": 6.0, "total_price": 12.0, "category": "groceries"},
        {"name": "Kawa", "quantity": 1, "unit_price": 9.5, "total_price": 9.5, "category": "cafe"},
        {"name": "Woda", "quantity": 1, "unit_price": 10.0, "total_price": 10.0, "category": "groceries"},
    ],
    "total_mismatch": True,
}

ERSTE_TXS = [
    {"date": "2026-06-16", "amount": -45.0, "currency": "PLN", "description": "ORLEN STACJA",
     "type": "expense", "raw_type": "OBCIĄŻENIE", "category": "transport", "category_display": "Транспорт"},
    {"date": "2026-06-16", "amount": -60.0, "currency": "PLN", "description": "",
     "type": "expense", "raw_type": "OBCIĄŻENIE", "category": "other", "category_display": "Другое",
     "foreign_card_no_merchant": True, "orig_amount": 14.0, "orig_currency": "EUR", "total_pln": 60.0},
    {"date": "2026-06-17", "amount": 2000.0, "currency": "PLN", "description": "Wynagrodzenie",
     "type": "income", "raw_type": "UZNANIE", "category": "other", "category_display": "Другое"},
]

REVOLUT_TXS = [
    {"date": "2026-06-14", "amount": 12.5, "currency": "EUR", "total_pln": 53.75, "description": "Uber",
     "type": "expense", "category": "transport", "category_display": "Транспорт"},
    {"date": "2026-06-15", "amount": 29.0, "currency": "PLN", "total_pln": 29.0, "description": "Netflix",
     "type": "expense", "category": "subscriptions", "category_display": "Подписки"},
]


# ── scenarios ────────────────────────────────────────────────────────────────

async def sc_start(c):
    await c.msg("/start")
    await c.msg("/help")
    await c.msg("/cancel")
    await c.cb("start:stats")
    await c.cb("start:budget")
    await c.cb("start:add")
    await c.msg("/cancel")


async def sc_reset(c):
    await c.msg("/reset")
    await c.msg("/reset confirm", user_id=RESET)


async def sc_add_flow(c):
    await c.msg("/add")
    await c.msg("abc")
    await c.msg("0 PLN")
    await c.msg("25.50 PLN")
    await c.msg("Żabka")
    await c.cb("cat:groceries")
    await c.msg("32.13.2026")
    await c.msg("19.06.2026")


async def sc_add_skip_and_budget_alert(c):
    await c.msg("/add")
    await c.msg("8 EUR")
    await c.msg("/skip")
    await c.cb("cat:cafe")
    await c.msg("/skip")


async def sc_add_cancel(c):
    await c.msg("/add")
    await c.msg("/cancel")
    await c.msg("/cancel")


async def sc_budget(c):
    await c.msg("/budget")
    await c.cb("budget:set")
    await c.cb("budget_set_cat:health")
    await c.msg("abc")
    await c.msg("200")
    await c.cb("budget_del:cafe")
    await c.cb("budget_del_yes:cafe")
    await c.cb("budget_del:electronics")
    await c.cb("budget_del_no")
    await c.cb("budget_show_all")
    await c.cb("budget:show")


async def sc_budget_empty(c):
    await c.msg("/budget", user_id=EMPTY)


async def sc_stats_categories(c):
    await c.msg("/stats")
    await c.cb("smenu:cats")
    await c.cb("speriod:cats:month")
    await c.cb("speriod:cats:week")
    await c.cb("speriod:cats:year")
    await c.cb("speriod:cats:all")


async def sc_stats_products(c):
    await c.cb("speriod:products:month")
    await c.cb("sprod:month:list")
    await c.cb("sproddet:month:0")
    await c.cb("sproddet:month:99")
    await c.cb("sprodback:month")
    await c.cb("sprodlist:month:0")
    await c.cb("speriod:products:day")
    await c.cb("sprod:day:list")
    await c.cb("sprod:month:search")
    await c.msg("coca")
    await c.cb("sprod:month:search")
    await c.msg("mle")
    await c.cb("sprod:month:search")
    await c.msg("zzzzqqq")
    await c.cb("sprod:year:search")
    await c.msg("l")
    await c.cb("sprod_pick:year:0")
    await c.cb("sprod_pick:year:99")


async def sc_stats_stores(c):
    await c.cb("speriod:stores:month")
    await c.cb("sstoredet:month:0")
    await c.cb("sstoredet:month:99")
    await c.cb("sstorelist:month")


async def sc_stats_legacy(c):
    await c.cb("stats:7")
    await c.cb("stats:30")
    await c.cb("stats:365")
    await c.cb("trend:30")
    await c.cb("stores:30")


async def sc_stats_compare(c):
    await c.cb("stats_compare")
    await c.cb("stats_compare:week")
    await c.cb("stats_compare:month")
    await c.cb("stats_compare:year")


async def sc_subscriptions(c):
    await c.msg("/subscriptions")
    await c.cb("subs:show")


async def sc_stats_empty(c):
    for data in ("speriod:cats:month", "speriod:products:week", "sprod:week:list", "speriod:stores:year",
                 "stats:7", "trend:7", "stores:7", "subs:show", "sproddet:month:0", "sstoredet:month:0",
                 "sprod_pick:month:0"):
        await c.cb(data, user_id=EMPTY)
    await c.cb("sprod:month:search", user_id=EMPTY)
    await c.msg("coca", user_id=EMPTY)


async def sc_export(c):
    await c.msg("/export")
    await c.cb("export_type:transactions")
    await c.cb("export_period:transactions:month")
    await c.cb("export_period:items:all")
    await c.cb("export_period:all:year")
    await c.cb("export_period:all:custom")
    await c.msg("bad")
    await c.msg("01.06.2026-20.06.2026")
    await c.cb("export_done:again")
    await c.cb("export_done:menu")
    await c.cb("export_period:transactions:today", user_id=EMPTY)


async def sc_reports_settings(c):
    await c.msg("/reports")
    await c.cb("report_toggle:daily")
    await c.cb("report_noop:daily")


async def sc_search(c):
    import bot.services.search_parser as sp
    from bot.services.llm import LLMUnavailableError

    def _unavailable():
        raise LLMUnavailableError("test: provider unavailable")
    c.monkeypatch.setattr(sp, "get_provider", _unavailable)
    await c.msg("/search")
    await c.msg("/search траты в Kaufland за июнь")
    await c.msg("/find покупки дороже 1000")
    await c.msg("/search дешевле 100")
    await c.cb("srchpage:1")
    await c.cb("srchpage:0", user_id=EMPTY)


async def sc_wrapped(c):
    await c.msg("/wrapped")
    await c.msg("/wrapped", user_id=EMPTY)


async def sc_split(c):
    rid = c.ids["kaufland"]
    await c.msg("/split")
    await c.cb(f"splitpick:{rid}")
    receipt = await __import__("bot.db.crud", fromlist=["x"]).get_receipt_by_id(c.session, rid)
    item_id = receipt.items[2].id
    await c.cb(f"splittog:{rid}:{item_id}:un")
    await c.cb(f"splitdone:{rid}")
    await c.cb(f"splittog:{rid}:{item_id}:on")
    await c.cb(f"splitdone:{rid}")
    await c.cb("splitpick:999999")
    await c.msg("/split", user_id=EMPTY)


async def sc_receipt_photo(c):
    _patch_receipt(c, parse_bank_transaction_screenshot=_async_return(None),
                   parse_receipt=_async_return(json.loads(json.dumps(PARSED_RECEIPT))))
    await c.msg(None, photo=True)
    rid = max(r.id for r in (await c.session.execute(
        __import__("sqlalchemy").select(Receipt))).scalars())
    await c.cb(f"recat:{rid}")
    await c.cb(f"recat_set:{rid}:cafe")
    await c.cb(f"recat_set:{rid}:bogus")
    await c.cb("recat_set:999999:cafe")


async def sc_receipt_bank_screenshot(c):
    _patch_receipt(c, parse_bank_transaction_screenshot=_async_return(
        {"merchant": "ORLEN", "date": "2026-06-20", "currency": "EUR", "amount": 250.0}))
    await c.msg(None, photo=True)


async def sc_receipt_errors(c):
    _patch_receipt(c, parse_bank_transaction_screenshot=_async_return(exc=RuntimeError("x")),
                   parse_receipt=_async_return(exc=ValueError("bad json")))
    await c.msg(None, photo=True)
    _patch_receipt(c, parse_receipt=_async_return(exc=RuntimeError("api down")))
    await c.msg(None, photo=True)


async def sc_documents(c):
    await c.msg(None, document={"file_name": "notes.txt", "mime_type": "text/plain"})
    await c.msg(None, document={"file_name": "scan.pdf", "mime_type": "application/pdf"})
    _patch_receipt(c, is_revolut_statement=lambda text: False)
    await c.msg(None, document={"file_name": "export.csv", "mime_type": "text/csv"})


async def sc_erste(c):
    _patch_receipt(c, is_erste_bank_statement=lambda b: True,
                   parse_erste_pdf=lambda b: json.loads(json.dumps(ERSTE_TXS)))
    await c.msg(None, document={"file_name": "wyciag.pdf", "mime_type": "application/pdf"})
    await c.cb("erste:save")
    await c.msg("Booking.com")
    await c.cb("erste:save")
    await c.msg(None, document={"file_name": "wyciag.pdf", "mime_type": "application/pdf"})
    await c.cb("erste:cancel")
    _patch_receipt(c, parse_erste_pdf=lambda b: [])
    await c.msg(None, document={"file_name": "wyciag.pdf", "mime_type": "application/pdf"})


async def sc_revolut(c):
    _patch_receipt(c, is_revolut_statement=lambda text: True,
                   parse_revolut_csv=lambda b: json.loads(json.dumps(REVOLUT_TXS)))
    await c.msg(None, document={"file_name": "revolut.csv", "mime_type": "text/csv"})
    await c.cb("revolut:save")
    await c.cb("revolut:save")
    await c.msg(None, document={"file_name": "revolut.csv", "mime_type": "text/csv"})
    await c.cb("revolut:cancel")


async def sc_reports_content(c):
    from bot.services import reports
    d = dt.date
    c.ret(await reports.build_daily_report(c.session, U, d(2026, 6, 19)))
    c.ret(await reports.build_daily_report(c.session, U, d(2026, 6, 17)))
    c.ret(await reports.build_weekly_report(c.session, U, d(2026, 6, 8), d(2026, 6, 14)))
    c.ret(await reports.build_weekly_report(c.session, U, d(2026, 1, 5), d(2026, 1, 11)))
    c.ret(await reports.build_monthly_report(c.session, U, 2026, 6))
    c.ret(await reports.build_monthly_report(c.session, U, 2026, 5))
    c.ret(await reports.build_monthly_report(c.session, U, 2025, 1))


async def sc_scheduler(c):
    """Scheduled jobs send through the bot; the session factory is swapped."""
    import contextlib

    import bot.scheduler as sched

    @contextlib.asynccontextmanager
    async def _session():
        yield c.session
    c.monkeypatch.setattr(sched, "get_session", _session)
    await sched._send_daily(c.bot)
    await sched._send_weekly(c.bot)
    await sched._send_monthly(c.bot)


async def sc_budget_services(c):
    from bot.services import budget
    c.ret(await budget.check_budget_alerts(c.session, U))
    c.ret(await budget.build_budget_text(c.session, U))
    c.ret(await budget.get_budget_status_text(c.session, U, "transport"))
    await budget.check_and_notify_budgets(c.session, U, c.bot)
    await budget.check_and_notify_budgets(c.session, U, c.bot)  # already notified: silent


async def sc_excel(c):
    from openpyxl import load_workbook

    from bot.services.export import build_excel
    wb = load_workbook(await build_excel(c.session, U))
    c.ret({ws.title: [[cell.value for cell in row] for row in ws.iter_rows(max_row=4)]
           for ws in wb.worksheets})


async def sc_formatters(c):
    from bot.utils import formatters as f
    c.ret([f.format_date_ru("2026-03-08"), f.format_date_ru(None), f.format_date_ru("garbage"),
           f.format_date(None), f.format_date(dt.date(2026, 12, 1)), f.format_month("2026-02"),
           f.format_category("household"), f.format_category("unknown")])


async def sc_keyboards(c):
    from bot.keyboards import inline as kb
    products = [{"normalized_name": "coca-cola", "total_spent": 50.0},
                {"normalized_name": "chleb", "total_spent": 20.0}]
    boards = {
        "stats_menu": kb.stats_menu_keyboard(),
        "compare_period": kb.compare_period_keyboard(),
        "new_period": kb.new_period_keyboard("cats"),
        "products_action": kb.products_action_keyboard("month"),
        "products_list": kb.products_list_keyboard(products, "month", 1, 3, 10),
        "product_detail": kb.product_detail_keyboard("month"),
        "stores_list": kb.stores_list_keyboard([{"store": "Lidl", "total_pln": 10.0}], "month"),
        "store_detail": kb.store_detail_keyboard("month"),
        "fuzzy": kb.fuzzy_matches_keyboard(["coca-cola", "cola zero"], "month"),
        "stats_period": kb.stats_period_keyboard(),
        "stats_extra": kb.stats_extra_keyboard(30),
        "recat": kb.recat_keyboard(1),
        "recat_categories": kb.recat_categories_keyboard(1),
        "budget_category": kb.budget_category_keyboard(),
        "budget_list": kb.budget_list_keyboard(["groceries", "cafe"]),
        "budget_delete_confirm": kb.budget_delete_confirm_keyboard("cafe"),
        "budget": kb.budget_keyboard(),
        "categories": kb.categories_keyboard(),
        "erste_save": kb.erste_save_keyboard(),
        "revolut_save": kb.revolut_save_keyboard(),
        "export_type": kb.export_type_keyboard(),
        "export_period": kb.export_period_keyboard("all"),
        "export_done": kb.export_done_keyboard(),
    }
    c.ret({name: h._keyboard_texts(markup) for name, markup in boards.items()})


async def sc_api(c):
    from bot.api.routes import budgets, stats, transactions
    c.ret(await transactions.recent_transactions(limit=20, user_id=U, db=c.session))
    c.ret(await stats.stats_by_category(period="month", user_id=U, db=c.session))
    c.ret(await budgets.get_budgets(user_id=U, db=c.session))


class FakeProvider:
    """Stands in for the LLM provider: answers by schema, so the real
    parse_receipt / validate_receipt / normalization code runs."""
    name = "fake"

    def __init__(self, receipt=None, screenshot=None, raise_exc=None):
        self.receipt, self.screenshot, self.raise_exc = receipt, screenshot, raise_exc
        self.images: list[int] = []

    async def generate_json(self, *, system, prompt, schema, tier, image=None, max_tokens=1024, schema_hint=None):
        from bot.services.normalization import NORMALIZATION_SCHEMA
        from bot.services.vision import BANK_TX_SCHEMA, RECEIPT_SCHEMA
        if self.raise_exc is not None:
            raise self.raise_exc
        if schema is BANK_TX_SCHEMA:
            return self.screenshot or {"is_transaction_screen": False}
        if schema is RECEIPT_SCHEMA:
            self.images.append(len(image) if isinstance(image, list) else 1)
            return json.loads(json.dumps(self.receipt))
        if schema is NORMALIZATION_SCHEMA:
            names = json.loads(prompt)
            return [{"original": n, "normalized": n.title(), "volume_ml": None} for n in names]
        raise AssertionError("unexpected schema")


def _use_provider(c, provider):
    import bot.services.normalization as norm
    import bot.services.vision as vision
    c.monkeypatch.setattr(vision, "get_provider", lambda: provider)
    c.monkeypatch.setattr(norm, "get_provider", lambda: provider)


RECEIPT_ITEMS = [
    {"name": "ChipsyCrun130-140g", "quantity": 1, "unit_price": 7.99, "total_price": 7.99,
     "category": "groceries", "deposit_type": None},
    {"name": "Rabat", "quantity": 1, "unit_price": -2.81, "total_price": -2.81,
     "category": "groceries", "deposit_type": None},
    {"name": "KawaRozpKronJa200g", "quantity": 1, "unit_price": 27.99, "total_price": 27.99,
     "category": "groceries", "deposit_type": None},
    {"name": "Puszka Kaucja", "quantity": 1, "unit_price": 0.5, "total_price": 0.5,
     "category": "other", "deposit_type": "wydanie"},
]


def _llm_receipt(total):
    return {"store": "Biedronka", "date": "2026-06-16", "currency": "PLN", "items": RECEIPT_ITEMS, "total": total}


async def sc_vision_totals(c):
    """Real parse_receipt + validator: Suma read / null total / mismatch."""
    for total in (33.67, None, 99.99):
        _use_provider(c, FakeProvider(receipt=_llm_receipt(total)))
        await c.msg(None, photo=True)


async def sc_vision_pdf_pages(c):
    import bot.handlers.receipt as rh
    provider = FakeProvider(receipt=_llm_receipt(None))
    _use_provider(c, provider)
    c.monkeypatch.setattr(rh, "_pdf_pages_to_jpeg", lambda b: [b"page1", b"page2"])
    c.monkeypatch.setattr(rh, "is_erste_bank_statement", lambda b: False)
    await c.msg(None, document={"file_name": "receipt.pdf", "mime_type": "application/pdf"})
    c.ret({"pages_sent_to_model": provider.images})


async def sc_vision_unavailable(c):
    from bot.services.llm import LLMUnavailableError
    import bot.handlers.receipt as rh
    _use_provider(c, FakeProvider(raise_exc=LLMUnavailableError("gemini: HTTP 429")))
    await c.msg(None, photo=True)
    c.monkeypatch.setattr(rh, "_pdf_pages_to_jpeg", lambda b: [b"page1"])
    c.monkeypatch.setattr(rh, "is_erste_bank_statement", lambda b: False)
    await c.msg(None, document={"file_name": "receipt.pdf", "mime_type": "application/pdf"})


async def sc_vision_bank_screenshot(c):
    _use_provider(c, FakeProvider(screenshot={
        "is_transaction_screen": True, "merchant": "Grande Bistro", "amount": 15.0, "currency": "PLN",
        "date": "2026-06-18", "year_visible": True, "raw_title": None}))
    await c.msg(None, photo=True)


SCENARIOS = {name[3:]: fn for name, fn in sorted(globals().items()) if name.startswith("sc_")}


# ── runner ───────────────────────────────────────────────────────────────────

async def run_all(db_session, monkeypatch) -> dict[str, list]:
    h.install_figure_spy(monkeypatch)
    bot = h.make_bot()
    dp, redis = h.get_dispatcher(fakeredis.aioredis.FakeRedis)
    results: dict[str, list] = {}
    for name, scenario in SCENARIOS.items():
        await _reset_db(db_session)
        await redis.flushall()
        await redis.set("fx_rate:EUR", "4.3")
        dp.fsm.storage.storage.clear()
        ids = await seed(db_session)
        h.RECORDER.take()
        with monkeypatch.context() as mp:
            await scenario(Ctx(db_session, redis, bot, dp, ids, mp))
        results[name] = h.RECORDER.take()
    return results


async def _reset_db(session):
    from bot.db.models import Base
    await session.rollback()
    for table in reversed(Base.metadata.sorted_tables):
        await session.execute(table.delete())
    await session.commit()


def _check(lang: str, results: dict[str, list]) -> None:
    path = SNAPSHOT_DIR / f"{lang}.json"
    rendered = json.dumps(results, ensure_ascii=False, indent=1, default=str) + "\n"
    if os.environ.get("UPDATE_SNAPSHOTS") or not path.exists():
        SNAPSHOT_DIR.mkdir(exist_ok=True)
        path.write_text(rendered, encoding="utf-8")
        if not os.environ.get("UPDATE_SNAPSHOTS"):
            pytest.fail(f"{path.name} did not exist and was created; review and commit it")
        return
    golden = json.loads(path.read_text(encoding="utf-8"))
    actual = json.loads(rendered)
    for name in sorted(set(golden) | set(actual)):
        assert actual.get(name) == golden.get(name), f"[{lang}] scenario {name!r} differs from snapshot"


@freeze_time(FROZEN_NOW)
async def test_ru_snapshots(db_session, monkeypatch):
    _check("ru", await run_all(db_session, monkeypatch))
