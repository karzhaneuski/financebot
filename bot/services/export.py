import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from bot.db import crud
from bot.i18n import _
from bot.markers import display_name
from bot.utils.formatters import format_category

_HEADER_FILL = PatternFill(start_color="AEC6CF", end_color="AEC6CF", fill_type="solid")
_HEADER_FONT = Font(bold=True)


def _apply_header(ws, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT


def _autowidth(ws) -> None:
    for col_cells in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col_cells), default=10)
        ws.column_dimensions[col_cells[0].column_letter].width = max_len + 3


async def build_excel(session: AsyncSession, user_id: int) -> io.BytesIO:
    receipts = await crud.get_receipts(session, user_id, days=90)

    wb = Workbook()

    # --- Sheet 1: Receipts ---
    ws1 = wb.active
    ws1.title = _("Receipts")
    _apply_header(ws1, [_("Date"), _("Store"), _("Total"), _("Currency"), _("Total PLN")])

    for r in receipts:
        ws1.append([
            r.date.strftime("%d.%m.%Y") if r.date else "",
            display_name(r.store) or "",
            float(r.total),
            r.currency,
            float(r.total_pln),
        ])

    _autowidth(ws1)

    # --- Sheet 2: Items ---
    ws2 = wb.create_sheet(_("Items"))
    _apply_header(ws2, [_("Date"), _("Store"), _("Item"), _("Qty"), _("Price"), _("Total"), _("Currency"),
                        _("Total PLN"), _("Category")])

    for r in receipts:
        date_str = r.date.strftime("%d.%m.%Y") if r.date else ""
        store_str = display_name(r.store) or ""
        for item in r.items:
            ws2.append([
                date_str,
                store_str,
                display_name(item.name),
                float(item.quantity),
                float(item.unit_price) if item.unit_price is not None else "",
                float(item.total_price),
                r.currency,
                round(r.to_pln(item.total_price), 2),
                format_category(item.category.value),
            ])

    _autowidth(ws2)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
